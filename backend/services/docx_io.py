"""DOCX import/export with track changes, comments, TOC, footnotes, watermark."""
from __future__ import annotations
from io import BytesIO
from typing import Any
import base64
import zipfile
from datetime import datetime
from lxml import etree
from docx import Document as PyDocx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn as docx_qn

from .ooxml_xml import (
    NS, q, make_insertion, make_deletion,
    build_comments_parts, build_footnotes_part, latex_to_omml,
)


# =========================================================================
# Export
# =========================================================================

def export_docx(model: dict[str, Any]) -> bytes:
    doc = PyDocx()

    meta = model.get("meta", {})
    doc.core_properties.title = meta.get("title", "")
    doc.core_properties.author = meta.get("author", "")

    _apply_page_setup(doc, model.get("pageSetup", {}))

    wm = model.get("watermark")
    if wm:
        _add_watermark(doc, wm)

    _apply_headers_footers(doc, model.get("headers", {}), model.get("footers", {}),
                           model.get("pageNumbers", False))

    for block in model.get("blocks", []) or []:
        _write_block(doc, block)

    revisions = model.get("revisions", []) or []
    if revisions:
        _inject_revisions(doc, revisions)

    buf = BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    if model.get("comments") or model.get("footnotes"):
        data = _attach_extension_parts(data, model)

    if model.get("_preserved", {}).get("parts"):
        data = _attach_preserved_parts(data, model["_preserved"]["parts"])

    return data


def _apply_page_setup(doc, setup):
    size_map = {"Letter": (8.5, 11), "A4": (8.27, 11.69), "A3": (11.69, 16.54),
                "A5": (5.83, 8.27), "Legal": (8.5, 14)}
    for section in doc.sections:
        name = setup.get("size", "Letter")
        if name in size_map:
            w, h = size_map[name]
            section.page_width = Inches(w)
            section.page_height = Inches(h)
        if setup.get("orientation") == "landscape":
            section.orientation = 1
            section.page_width, section.page_height = section.page_height, section.page_width
        for key, attr in [("marginTop", "top_margin"), ("marginBottom", "bottom_margin"),
                          ("marginLeft", "left_margin"), ("marginRight", "right_margin")]:
            if setup.get(key) is not None:
                setattr(section, attr, Inches(float(setup[key]) / 96))


def _write_block(doc, block):
    btype = block.get("type", "paragraph")
    attrs = block.get("attrs", {}) or {}

    if btype == "paragraph":
        style = attrs.get("style", "Normal")
        p = doc.add_paragraph(style=style if style != "Normal" else None)
        for run in block.get("runs", []):
            r = p.add_run(run.get("text", ""))
            _apply_run_fmt(r, run.get("attrs", {}))

    elif btype == "heading":
        level = int(attrs.get("level", 1))
        p = doc.add_heading(level=level)
        for run in block.get("runs", []):
            r = p.add_run(run.get("text", ""))
            _apply_run_fmt(r, run.get("attrs", {}))

    elif btype == "list":
        style = "List Bullet" if attrs.get("ordered") is False else "List Number"
        for item in block.get("items", []):
            p = doc.add_paragraph(style=style)
            for run in item:
                r = p.add_run(run.get("text", ""))
                _apply_run_fmt(r, run.get("attrs", {}))

    elif btype == "table":
        rows = block.get("rows", [])
        if not rows:
            return
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                t.cell(i, j).text = str(cell)

    elif btype == "image":
        src = attrs.get("src", "")
        if src.startswith("data:image/"):
            try:
                _, b64 = src.split(",", 1)
                doc.add_picture(BytesIO(base64.b64decode(b64)), width=Inches(attrs.get("width", 4)))
            except Exception:
                pass

    elif btype == "pagebreak":
        doc.add_page_break()

    elif btype == "toc":
        _insert_toc_field(doc, attrs.get("levels", "1-3"))

    elif btype == "field":
        _insert_field(doc, attrs.get("name", "PAGE"))

    elif btype == "equation":
        _insert_equation(doc, attrs.get("latex", ""))

    elif btype == "footnote":
        _insert_footnote_ref(doc, attrs.get("id", 1))

    elif btype == "caption":
        p = doc.add_paragraph()
        r = p.add_run(f"Figure {attrs.get('n', 1)}: ")
        r.bold = True
        r2 = p.add_run(attrs.get("text", ""))
        r2.italic = True


def _apply_run_fmt(run, attrs):
    if attrs.get("bold"): run.bold = True
    if attrs.get("italic"): run.italic = True
    if attrs.get("underline"): run.underline = True
    if attrs.get("strike"): run.font.strike = True
    if attrs.get("sub"): run.font.subscript = True
    if attrs.get("sup"): run.font.superscript = True
    if attrs.get("size"):
        try: run.font.size = Pt(float(attrs["size"]))
        except Exception: pass
    if attrs.get("font"): run.font.name = attrs["font"]
    if attrs.get("color"):
        try:
            run.font.color.rgb = RGBColor.from_string(attrs["color"].lstrip("#").upper())
        except Exception: pass


def _insert_toc_field(doc, levels="1-3"):
    p = doc.add_paragraph()
    r = p.add_run()
    r._element.append(_field_char("begin"))
    r._element.append(_field_instr(f' TOC \\o "{levels}" \\h \\z \\u '))
    r._element.append(_field_char("separate"))
    t = etree.SubElement(r._element, docx_qn("w:t"))
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = "Right-click and choose Update Field to build the table of contents."
    r._element.append(_field_char("end"))


def _insert_field(doc, name: str):
    p = doc.add_paragraph()
    r = p.add_run()
    r._element.append(_field_char("begin"))
    r._element.append(_field_instr(f" {name} "))
    r._element.append(_field_char("separate"))
    t = etree.SubElement(r._element, docx_qn("w:t"))
    t.text = "1"
    r._element.append(_field_char("end"))


def _insert_equation(doc, latex: str):
    p = doc.add_paragraph()
    omml = latex_to_omml(latex)
    omml_str = etree.tostring(omml)
    omml_str = omml_str.replace(
        b"<m:oMath",
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"',
        1,
    )
    try:
        p._element.append(etree.fromstring(omml_str))
    except Exception:
        r = p.add_run(latex)
        r.italic = True


def _insert_footnote_ref(doc, footnote_id: int):
    p = doc.add_paragraph()
    r = p.add_run()
    rpr = etree.SubElement(r._element, docx_qn("w:rPr"))
    style = etree.SubElement(rpr, docx_qn("w:rStyle"))
    style.set(docx_qn("w:val"), "FootnoteReference")
    ref = etree.SubElement(r._element, docx_qn("w:footnoteReference"))
    ref.set(docx_qn("w:id"), str(footnote_id))


def _field_char(kind: str) -> etree._Element:
    el = etree.Element(docx_qn("w:fldChar"))
    el.set(docx_qn("w:fldCharType"), kind)
    return el


def _field_instr(text: str) -> etree._Element:
    el = etree.Element(docx_qn("w:instrText"))
    el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    el.text = text
    return el


def _add_watermark(doc, wm: dict):
    text = wm.get("text", "DRAFT")
    color = wm.get("color", "#C0C0C0")
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = 1
        run = p.add_run()
        vml = f'''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="PowerPlusWaterMarkObject" style="position:absolute;left:0;text-align:left;margin-left:0;margin-top:0;width:400pt;height:200pt;rotation:315;z-index:-251658752;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" fillcolor="{color}" stroked="f" xmlns:o="urn:schemas-microsoft-com:office:office"><v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{text}"/></v:shape></w:pict>'''
        try:
            run._element.append(etree.fromstring(vml))
        except Exception:
            pass


def _apply_headers_footers(doc, headers: dict, footers: dict, page_numbers: bool):
    for section in doc.sections:
        if headers.get("default"):
            _set_header_footer_text(section.header, headers["default"])
        if headers.get("first"):
            section.different_first_page_header_footer = True
            _set_header_footer_text(section.first_page_header, headers["first"])
        if footers.get("default"):
            _set_header_footer_text(section.footer, footers["default"])
        if page_numbers:
            _add_page_number_field(section.footer)


def _set_header_footer_text(container, text: str):
    p = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
    for r in list(p.runs):
        r.text = ""
    p.text = text


def _add_page_number_field(container):
    p = container.add_paragraph()
    p.alignment = 1
    r = p.add_run()
    r._element.append(_field_char("begin"))
    r._element.append(_field_instr(" PAGE "))
    r._element.append(_field_char("separate"))
    t = etree.SubElement(r._element, docx_qn("w:t"))
    t.text = "1"
    r._element.append(_field_char("end"))


def _inject_revisions(doc, revisions):
    for rev in revisions:
        p = doc.add_paragraph()
        author = rev.get("author", "User")
        when = datetime.fromisoformat(rev.get("date", datetime.utcnow().isoformat()))
        p._element.append(make_insertion(author, when, [{"text": rev.get("text", "")}]))
        if rev.get("deleted"):
            p._element.append(make_deletion(author, when, [{"text": rev.get("deleted")}]))


def _attach_extension_parts(docx_bytes: bytes, model: dict) -> bytes:
    zin = zipfile.ZipFile(BytesIO(docx_bytes))
    zout_buf = BytesIO()
    with zipfile.ZipFile(zout_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))

        if model.get("comments"):
            parts = build_comments_parts(model["comments"])
            for name, content in parts.items():
                zout.writestr(name, content)
            ct = zin.read("[Content_Types].xml")
            ct = _ensure_content_types(ct, [
                ("application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml", "/word/comments.xml"),
                ("application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml", "/word/commentsExtended.xml"),
                ("application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml", "/word/people.xml"),
            ])
            zout.writestr("[Content_Types].xml", ct)

            rels = zin.read("word/_rels/document.xml.rels")
            rels = _ensure_rels(rels, [
                ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", "comments.xml"),
                ("http://schemas.microsoft.com/office/2011/relationships/commentsExtended", "commentsExtended.xml"),
                ("http://schemas.microsoft.com/office/2011/relationships/people", "people.xml"),
            ])
            zout.writestr("word/_rels/document.xml.rels", rels)

        if model.get("footnotes"):
            zout.writestr("word/footnotes.xml", build_footnotes_part(model["footnotes"]))
            ct = zout.getvalue() if False else None

    return zout_buf.getvalue()


def _ensure_content_types(ct_bytes: bytes, entries: list[tuple[str, str]]) -> bytes:
    root = etree.fromstring(ct_bytes)
    existing_parts = {el.get("PartName") for el in root.findall(q("ct", "Override"))}
    for ctype, part in entries:
        if part in existing_parts:
            continue
        el = etree.SubElement(root, q("ct", "Override"))
        el.set("ContentType", ctype)
        el.set("PartName", part)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _ensure_rels(rels_bytes: bytes, entries: list[tuple[str, str]]) -> bytes:
    root = etree.fromstring(rels_bytes)
    existing_targets = {el.get("Target") for el in root.findall(q("pr", "Relationship"))}
    used_ids = {el.get("Id") for el in root.findall(q("pr", "Relationship"))}
    for rtype, target in entries:
        if target in existing_targets:
            continue
        next_n = max([int(i[3:]) for i in used_ids if i.startswith("rId") and i[3:].isdigit()] or [0]) + 1
        rid = "rId" + str(next_n)
        el = etree.SubElement(root, q("pr", "Relationship"))
        el.set("Id", rid)
        el.set("Type", rtype)
        el.set("Target", target)
        used_ids.add(rid)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _attach_preserved_parts(docx_bytes: bytes, parts: dict[str, str]) -> bytes:
    zin = zipfile.ZipFile(BytesIO(docx_bytes))
    zout_buf = BytesIO()
    with zipfile.ZipFile(zout_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        for name, content_b64 in parts.items():
            zout.writestr(name, base64.b64decode(content_b64))
    return zout_buf.getvalue()


# =========================================================================
# Import
# =========================================================================

def import_docx(data: bytes) -> dict[str, Any]:
    doc = PyDocx(BytesIO(data))
    blocks = []
    for p in doc.paragraphs:
        style_name = (p.style.name or "Normal") if p.style else "Normal"
        runs = [{"text": r.text, "attrs": _extract_run_attrs(r)} for r in p.runs]
        if not runs and p.text:
            runs = [{"text": p.text, "attrs": {}}]
        if style_name.lower().startswith("heading"):
            try:
                level = int(style_name.split()[-1])
            except Exception:
                level = 1
            blocks.append({"type": "heading", "runs": runs, "attrs": {"level": level}})
        else:
            blocks.append({"type": "paragraph", "runs": runs, "attrs": {"style": style_name}})

    for t in doc.tables:
        rows = [[c.text for c in row.cells] for row in t.rows]
        blocks.append({"type": "table", "rows": rows})

    preserved = {}
    try:
        zf = zipfile.ZipFile(BytesIO(data))
        for name in zf.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                continue
            if any(name.startswith(p) for p in ("customXml/", "docProps/custom", "word/embeddings/")):
                preserved[name] = base64.b64encode(zf.read(name)).decode("ascii")
    except Exception:
        pass

    return {
        "meta": {"title": doc.core_properties.title or "", "author": doc.core_properties.author or ""},
        "pageSetup": {"size": "Letter", "orientation": "portrait",
                      "marginTop": 96, "marginBottom": 96, "marginLeft": 96, "marginRight": 96},
        "styles": {},
        "blocks": blocks,
        "headers": {}, "footers": {}, "comments": [], "revisions": [],
        "footnotes": [],
        "schemaVersion": 4,
        "_preserved": {"parts": preserved},
    }


def _extract_run_attrs(run) -> dict:
    a: dict = {}
    if run.bold: a["bold"] = True
    if run.italic: a["italic"] = True
    if run.underline: a["underline"] = True
    if run.font.strike: a["strike"] = True
    if run.font.size:
        try: a["size"] = run.font.size.pt
        except Exception: pass
    if run.font.name: a["font"] = run.font.name
    try:
        if run.font.color and run.font.color.rgb:
            a["color"] = "#" + str(run.font.color.rgb)
    except Exception:
        pass
    return a